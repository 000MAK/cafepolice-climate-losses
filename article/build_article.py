"""Build the 1.5-page Café Poliçe article (Turkish) as .docx and .md from outputs/metrics.json."""
from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
M = json.loads((ROOT / "outputs" / "metrics.json").read_text(encoding="utf-8"))
NAVY, ORANGE = RGBColor(0x14, 0x21, 0x3D), RGBColor(0xB4, 0x56, 0x1C)
IAN_NCEI_BN = 112.9   # NOAA NCEI Billion-Dollar Disasters, Hurricane Ian (2022)


def tr(x: float, d: int = 0) -> str:
    """Turkish number format: thousands '.', decimals ','."""
    s = f"{x:,.{d}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def pct(x: float, d: int = 1) -> str:
    return "%" + tr(x * 100, d)


h, s, dm, p, ho = M["large_hail"], M["large_hail_share"], M["large_damage_episodes"], M["projection_2030"], M["holdout"]
ht, dt = h["trend"], dm["trend"]

TITLE = "Dolu Büyüyor, Veri Küçülüyor"
SUBTITLE = "Otuz yıllık açık hava verisiyle iklim hasarını geleceğe taşımak"
BYLINE = "Mehmet Ali Kaya · Veri Yönetimi Uzmanı, Yönetim Bilişim Sistemleri Doktora Adayı, İstatistik Yüksek Lisans"
SPOT = (f"ABD'nin otuz yıllık fırtına kayıtlarında 5 santimetreyi aşan dolu her yıl {pct(ht['annual_change'])} artıyor; "
        "aynı veri setinin hasar sütunu ise riskin azaldığını söylüyor. Geleceği fiyatlamak isteyen sigortacı, "
        "önce hangi sütuna güveneceğine karar vermek zorunda.")

SECTIONS = [
    ("Veri ve problem", [
        f"Kullandığım veri, ABD Ulusal Okyanus ve Atmosfer İdaresi'nin (NOAA) kamuya açık Storm Events veritabanı: 1996–2025 arasında kaydedilmiş "
        f"{tr(M['events'])} hava olayı, bunların {tr(M['hail_reports'])}'i dolu raporu. Hasar tutarlarını tüketici fiyat endeksiyle 2025 dolarına taşıdım. "
        "Sigortacı için soru basit: Geçmişin ortalamasıyla kurulan bir tarife, 2030'un dolusunu karşılar mı?",
        f"İlk bulgu bir uyarı oldu. Veri setinde 10 milyon doları aşan hasar kaydı olan fırtına sayısı yılda {pct(-dt['annual_change'])} azalıyor; "
        f"1996–2005'te yılda {tr(dm['mean_1996_2005'])}, son on yılda {tr(dm['mean_2016_2025'])}. Oysa olayların hasar tahmini taşıma oranı "
        f"{pct(M['damage_nonzero_share']['2008_2011'], 0)}'den {pct(M['damage_nonzero_share']['2022_2025'], 0)}'ye düşmüş. Ian Kasırgası bu kayıtta "
        f"{tr(M['largest_2022_episode_bn'], 1)} milyar dolar görünüyor; NOAA'nın kendi milyar dolarlık afetler listesi aynı fırtınayı {tr(IAN_NCEI_BN, 1)} milyar dolarla anıyor. "
        "Riskin azaldığı yok; azalan, hasarın kayda geçirilme oranı.",
    ]),
    ("Model ve bulgular", [
        f"Bu yüzden hasar sütununu bırakıp fiziksel tehlikeye, 5 santimetre (2 inç) ve üzeri dolu raporlarına baktım. Bu eşik, ABD'nin 2010'da şiddetli dolu "
        f"kriterini değiştirmesinden etkilenmiyor. Yıl trendli bir Poisson modeli büyük dolunun yılda {pct(ht['annual_change'])} arttığını gösteriyor "
        f"(güven aralığı {pct(ht['annual_change_ci95'][0])}–{pct(ht['annual_change_ci95'][1])}). Rapor sayısındaki genel artışa karşı da sağlam: "
        f"Tüm dolu raporları içinde büyük dolunun payı {pct(s['2010_2014'])}'dan {pct(s['2021_2025'])}'e çıkmış.",
        f"Modeli zamana karşı sınadım: 1996–2015 verisiyle kurup 2016–2025'i tahmin ettirdim. Trend modeli on yıl için {tr(ho['predicted_total_trend'])} büyük dolu raporu "
        f"bekledi; gerçekleşen {tr(ho['observed_total'])}. \"Geçmiş tekrar eder\" diyen düz ortalama ise {tr(ho['predicted_total_flat'])} bekleyip yüzde "
        f"{tr((1 - ho['predicted_total_flat'] / ho['observed_total']) * 100)} ıskaladı. 2030 için model yılda ortalama {tr(p['mean'])} büyük dolu raporu öngörüyor "
        f"(yüzde 80 aralık {tr(p['pi80'][0])}–{tr(p['pi80'][1])}); bu, otuz yıllık ortalamanın yüzde {tr(p['vs_30yr_mean'] * 100)} üzerinde.",
    ]),
    ("Geleceğin modelleri", [
        "Klasik istatistik, yani bu yazıdaki Poisson trend modeli gibi GLM'ler, açıklanabilir ve denetlenebilir olduğu için tarifenin omurgası olmaya devam edecek. "
        "Makine öğrenmesi, poliçe ve bina özelliklerinden riski sıralamada GLM'i geçiyor; derin öğrenme ise radar ve uydu görüntülerinden dolu izini "
        "çıkarmada güçlü, ama büyük ve etiketli veri istiyor. Bu çalışmanın asıl dersi ise gri kutu yaklaşımdan geliyor: Fiziksel tehlike ölçüsü "
        "(dolunun boyu) ile istatistiksel trendi birleştiren model, hasar kaydı çöktüğünde bile sinyali korudu. Gelecek, tek bir yöntemin değil; "
        "fiziksel veri, istatistik ve öğrenen modellerin birbirini denetlediği hibrit yapıların.",
    ]),
]

TABLE = [("Gösterge", "1996–2005", "2016–2025", "Model / test"),
         ("5 cm üstü dolu raporu (yıllık)", tr(h["mean_1996_2005"]), tr(h["mean_2016_2025"]), f"yılda +{pct(ht['annual_change'])}"),
         ("Büyük dolunun tüm dolu içindeki payı", f"{pct(s['2010_2014'])} (2010–14)", f"{pct(s['2021_2025'])} (2021–25)", "neredeyse iki kat"),
         ("10 milyon $ üstü kayıtlı hasar (yıllık)", tr(dm["mean_1996_2005"]), tr(dm["mean_2016_2025"]), f"yılda −{pct(-dt['annual_change'])}"),
         ("2016–2025 büyük dolu toplamı", "—", f"{tr(ho['observed_total'])} gerçekleşen", f"trend {tr(ho['predicted_total_trend'])} · düz {tr(ho['predicted_total_flat'])}")]

LESSON = (f"Geçmiş ortalamayla kalibre edilen bir dolu tarifesi, bu gidişle 2030'da tehlikeyi yüzde {tr((1 - h['mean_1996_2025'] / p['mean']) * 100)} eksik görecek. Üstelik hasar verisine körü körüne "
          "güvenen bir model, risk artarken \"azalıyor\" sonucuna varabilir. Türkiye'de dolu kasko, konut ve tarım portföylerinin ortak derdi; bu yüzden "
          "hasar kayıtlarının yanında meteorolojik tehlike verisini de tarifeye bağlamak, sigortacılığın geleceği için bir tercih değil, zorunluluk.")
ACTIONS = [
    "Hasar verisini tehlike verisiyle eşleştirin. Poliçe ve hasar dosyalarını, olay günü ve konumuyla meteorolojik kayıtlara bağlamak, "
    "hasar kaydı eksik kaldığında bile riskin yönünü görmeyi sağlar.",
    "Tarifeyi ortalamaya değil, trende kalibre edin. Otuz yıllık ortalama, artan bir tehlikede her yıl biraz daha geride kalır; "
    "trend katsayısını ve belirsizliğini fiyatlama notuna yazmak bu açığı görünür kılar.",
    "Modeli her yıl geleceğe karşı sınayın. Geçmişin bir bölümünü saklayıp modelin onu tahmin edip edemediğine bakmak, "
    "tek bir başarı ölçüsünden çok daha dürüst bir kalite kontrolüdür.",
]
GITHUB = "https://github.com/000MAK/cafepolice-climate-losses"
HF_DATA = "https://huggingface.co/datasets/makactuary/noaa-storm-events-hail-1996-2025"
HF_MODEL = "https://huggingface.co/makactuary/noaa-hail-trend-2030"
LIMITS = ("Sınırlılıklar: Veriler ABD'ye ait ve sigortalı hasar değil, ekonomik hasar tahminidir. Dolu raporları gözlemci ve nüfus yoğunluğundan etkilenir; "
          "pay analizi bu etkiyi azaltır ama yok etmez. Trendin ne kadarının iklimden, ne kadarının raporlama ve yapılaşmadan geldiği bu veriyle ayrıştırılamaz.")
SOURCES = ("Kaynaklar: NOAA NCEI Storm Events Database (1996–2025, kamu malı); NOAA NCEI Billion-Dollar Weather and Climate Disasters; "
           "ABD tüketici fiyat endeksi (FRED, CPIAUCSL).")


def bottom_border(paragraph, color="E07A2E") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "4"), ("w:color", color)):
        bottom.set(qn(k), v)
    borders.append(bottom)
    p_pr.append(borders)


def build_docx(path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    hs = doc.styles["Heading 1"]
    hs.font.name, hs.font.size, hs.font.color.rgb, hs.font.bold = "Calibri", Pt(13), NAVY, True
    hs.paragraph_format.space_before, hs.paragraph_format.space_after = Pt(10), Pt(4)

    for text, size, color, bold, italic in ((TITLE, 24, NAVY, True, False), (SUBTITLE, 12.5, ORANGE, False, False),
                                            (BYLINE, 9.5, None, False, True), (SPOT, 11.5, None, True, False)):
        para = doc.add_paragraph()
        r = para.add_run(text)
        r.bold, r.italic, r.font.size = bold, italic, Pt(size)
        if color:
            r.font.color.rgb = color
        if text == BYLINE:
            bottom_border(para)

    for heading, paras in SECTIONS:
        doc.add_heading(heading, level=1)
        for text in paras:
            doc.add_paragraph(text)
        if heading == "Model ve bulgular":
            table = doc.add_table(rows=len(TABLE), cols=4)
            table.style = "Light Grid Accent 1"
            for i, vals in enumerate(TABLE):
                for j, val in enumerate(vals):
                    cell = table.cell(i, j)
                    cell.text = val
                    run = cell.paragraphs[0].runs[0]
                    run.font.size, run.bold = Pt(9), i == 0
            cp = doc.add_paragraph()
            r = cp.add_run("Tablo 1. NOAA Storm Events, 1996–2025. Hasar tutarları 2025 doları; test: 1996–2015 ile kurulan modelin 2016–2025 tahmini.")
            r.italic, r.font.size = True, Pt(8.5)
            doc.add_picture(str(ROOT / "outputs" / "charts" / "01_dolu_ve_hasar_kaydi.png"), width=Cm(15.4))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cp.add_run("Grafik 1. Solda büyük dolu raporları, trend modeli ve 2030 tahmini; sağda tehlike ile kayıtlı büyük hasarın endeksi.")
            r.italic, r.font.size = True, Pt(9)

    doc.add_heading("Sektör için tek ders", level=1)
    doc.add_paragraph(LESSON)
    doc.add_heading("Yarın sabah ne yapmalı?", level=1)
    for i, text in enumerate(ACTIONS, 1):
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Cm(0.5)
        para.add_run(f"{i}. ").bold = True
        para.add_run(text)
    doc.add_heading("Kod, veri ve model", level=1)
    for label, url in (("Kod ve sonuçlar (GitHub): ", GITHUB), ("Veri (Hugging Face): ", HF_DATA), ("Model ve 2030 tahmini (Hugging Face): ", HF_MODEL)):
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(2)
        para.add_run(label).bold = True
        para.add_run(url)
    para = doc.add_paragraph()
    r = para.add_run(f"{LIMITS} {SOURCES}")
    r.italic, r.font.size = True, Pt(8.5)
    doc.save(path)


def build_md(path: Path) -> None:
    lines = [f"# {TITLE}", f"*{SUBTITLE}*", "", f"*{BYLINE}*", "", f"**{SPOT}**", ""]
    for heading, paras in SECTIONS:
        lines += [f"## {heading}", ""] + [x + "\n" for x in paras]
        if heading == "Model ve bulgular":
            lines += ["| " + " | ".join(TABLE[0]) + " |", "|---|---|---|---|"] + ["| " + " | ".join(r) + " |" for r in TABLE[1:]]
            lines += ["", "![Grafik 1](../outputs/charts/01_dolu_ve_hasar_kaydi.png)", ""]
    lines += ["## Sektör için tek ders", "", LESSON, "", "## Yarın sabah ne yapmalı?", ""] + [f"{i}. {a}" for i, a in enumerate(ACTIONS, 1)]
    lines += ["", "## Kod, veri ve model", "", f"- Kod: {GITHUB}", f"- Veri: {HF_DATA}", f"- Model: {HF_MODEL}", "", f"*{LIMITS}*", "", f"*{SOURCES}*", ""]
    path.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    out = ROOT / "article"
    name = "CafePolice_Dolu_IklimHasari_MehmetAliKaya"
    build_docx(out / f"{name}.docx")
    build_md(out / f"{name}.md")
    if len(sys.argv) > 1:
        build_docx(Path(sys.argv[1]) / f"{name}.docx")
    print("written")
